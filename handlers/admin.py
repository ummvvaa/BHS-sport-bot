"""Админские команды: /stats, /analytics, /remind, /export, /reload. Доступны только ADMIN_IDS."""
from __future__ import annotations

import asyncio
import html
import io
import logging
from datetime import datetime

from aiogram import F, Router
from aiogram.exceptions import TelegramBadRequest, TelegramForbiddenError
from aiogram.filters import Command, CommandObject
from aiogram.types import BufferedInputFile, CallbackQuery, Message
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

import sheets
from config import settings
from keyboards import CB_REMIND_NO, CB_REMIND_YES, remind_confirm_kb
from sheets import CLASSES, DATE_FORMAT, GROUPS, TZ, Analytics, Enrollment, ExportData, SectionLimit, Stats
from texts import days_ru, t

logger = logging.getLogger(__name__)
router = Router(name="admin")
router.message.filter(F.from_user.id.in_(settings.admin_ids))
router.callback_query.filter(F.from_user.id.in_(settings.admin_ids))

LANG = "ru"
SUMMARY_SHEET = "Сводка"
FORBIDDEN_SHEET_CHARS = set('[]:*?/\\')
REMIND_DELAY_SECONDS = 0.05
MAX_ANALYTICS_DAYS = 365


# --------------------------------------------------------------------------- #
# /export
# --------------------------------------------------------------------------- #


def _safe_sheet_title(title: str, used: set[str]) -> str:
    cleaned = "".join("_" if ch in FORBIDDEN_SHEET_CHARS else ch for ch in title).strip() or "Секция"
    cleaned = cleaned[:31]
    candidate = cleaned
    counter = 2
    while candidate in used:
        suffix = f" ({counter})"
        candidate = cleaned[: 31 - len(suffix)] + suffix
        counter += 1
    used.add(candidate)
    return candidate


def _write_table(ws: Worksheet, headers: list[str], rows: list[list[object]]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    for col_index, header in enumerate(headers, start=1):
        width = len(str(header))
        for row in rows:
            value = row[col_index - 1] if col_index - 1 < len(row) else ""
            width = max(width, len(str(value)))
        ws.column_dimensions[get_column_letter(col_index)].width = min(width + 2, 50)


def _section_rows(items: list[Enrollment]) -> list[list[object]]:
    return [
        [index, item.name, item.class_num, item.phone, item.enrolled_at]
        for index, item in enumerate(items, start=1)
    ]


def _summary_rows(stats: Stats) -> list[list[object]]:
    rows: list[list[object]] = []
    for section in stats.sections:
        row: list[object] = [section.name, section.count_total]
        for group in GROUPS:
            row.append(f"{section.count_by_group.get(group, 0)}/{section.limit.group_limit(group)}")
        rows.append(row)
    return rows


def build_xlsx(data: ExportData) -> bytes:
    wb = Workbook()
    summary = wb.active
    summary.title = SUMMARY_SHEET
    _write_table(summary, ["Секция", "Записано", *GROUPS], _summary_rows(data.stats))

    used: set[str] = {SUMMARY_SHEET}
    limit: SectionLimit
    for limit, items in data.sections:
        ws = wb.create_sheet(title=_safe_sheet_title(limit.name, used))
        _write_table(ws, ["№", "Имя", "Класс", "Телефон", "Дата записи"], _section_rows(items))

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@router.message(Command("export"))
async def cmd_export(message: Message) -> None:
    data = await sheets.get_export_data()
    if not data.sections:
        await message.answer(t(LANG, "admin_no_sections"))
        return
    now = datetime.now(TZ)
    content = build_xlsx(data)
    filename = f"sections_{now.strftime('%Y-%m-%d_%H-%M')}.xlsx"
    await message.answer_document(
        BufferedInputFile(content, filename=filename),
        caption=t(LANG, "export_caption", when=now.strftime(DATE_FORMAT)),
    )
    logger.info("Админ %s выгрузил %s", message.from_user.id, filename)


# --------------------------------------------------------------------------- #
# /stats
# --------------------------------------------------------------------------- #


@router.message(Command("stats"))
async def cmd_stats(message: Message) -> None:
    stats = await sheets.get_stats()
    lines: list[str] = []
    for section in stats.sections:
        by_group = ", ".join(
            t(
                LANG,
                "stats_group",
                group=group,
                count=section.count_by_group.get(group, 0),
                limit=section.limit.group_limit(group),
            )
            for group in GROUPS
        )
        lines.append(
            t(LANG, "stats_line", section=section.name, count=section.count_total, limit=section.limit.total, by_group=by_group)
        )
    if not lines:
        lines.append(t(LANG, "admin_no_sections"))
    lines.append("")
    lines.append(t(LANG, "stats_total", registered=stats.registered, enrolled=stats.enrolled))
    if stats.deadline is None:
        lines.append(t(LANG, "stats_no_deadline"))
    else:
        lines.append(t(LANG, "stats_deadline", deadline=stats.deadline.strftime(DATE_FORMAT)))
    lines.append(t(LANG, "stats_closed") if stats.closed else t(LANG, "stats_open"))
    await message.answer("\n".join(lines))


# --------------------------------------------------------------------------- #
# /analytics
# --------------------------------------------------------------------------- #


def _mono_table(headers: list[str], rows: list[list[str]]) -> str:
    """Таблица с выровненными колонками для <pre>."""
    widths = [len(h) for h in headers]
    for row in rows:
        for index, cell in enumerate(row):
            widths[index] = max(widths[index], len(cell))
    lines = [" | ".join(cell.ljust(widths[i]) for i, cell in enumerate(headers))]
    for row in rows:
        lines.append(" | ".join(cell.ljust(widths[i]) for i, cell in enumerate(row)))
    return "\n".join(line.rstrip() for line in lines)


def build_analytics_text(analytics: Analytics) -> str:
    title = (
        t(LANG, "analytics_title_days", period=days_ru(analytics.days))
        if analytics.days
        else t(LANG, "analytics_title_all")
    )
    rows = [
        [
            str(class_num),
            str(analytics.per_class[class_num].registered),
            str(analytics.per_class[class_num].enrolled),
            str(analytics.per_class[class_num].without_section),
        ]
        for class_num in CLASSES
    ]
    table = _mono_table(["Класс", "Зарегались", "Записались", "Без секции"], rows)

    lines = [title, "", f"<pre>{html.escape(table)}</pre>", ""]
    passive = analytics.most_passive
    if passive is not None:
        lines.append(t(LANG, "analytics_passive", class_num=passive[0], count=passive[1]))
    lines.append(t(LANG, "analytics_total", registered=analytics.registered, enrolled=analytics.enrolled))
    lines.append("")
    if analytics.by_day:
        lines.append(t(LANG, "analytics_by_day"))
        if analytics.by_day_truncated:
            lines.append(t(LANG, "analytics_by_day_truncated", n=len(analytics.by_day)))
        for day, counts in analytics.by_day:
            parts = ", ".join(f"{class_num}кл {counts.get(class_num, 0)}" for class_num in CLASSES)
            lines.append(f"{day}: {parts}")
    else:
        lines.append(t(LANG, "analytics_by_day_empty"))
    return "\n".join(lines)


def _parse_days(args: str | None) -> int | None | ValueError:
    text = (args or "").strip()
    if not text:
        return None
    if not text.isdigit() or not 1 <= int(text) <= MAX_ANALYTICS_DAYS:
        return ValueError(text)
    return int(text)


@router.message(Command("analytics"))
async def cmd_analytics(message: Message, command: CommandObject) -> None:
    days = _parse_days(command.args)
    if isinstance(days, ValueError):
        await message.answer(t(LANG, "analytics_usage"))
        return
    analytics = await sheets.get_analytics(days)
    await message.answer(build_analytics_text(analytics), parse_mode="HTML")


# --------------------------------------------------------------------------- #
# /remind
# --------------------------------------------------------------------------- #


def _parse_remind_target(args: str | None) -> int | None | ValueError:
    """«all» → None (все классы), «8..11» → номер класса, иначе ValueError."""
    text = (args or "").strip().lower()
    if text == "all":
        return None
    if text.isdigit() and int(text) in CLASSES:
        return int(text)
    return ValueError(text)


def _target_label(class_num: int | None) -> str:
    if class_num is None:
        return t(LANG, "remind_target_all")
    return t(LANG, "remind_target_class", class_num=class_num)


@router.message(Command("remind"))
async def cmd_remind(message: Message, command: CommandObject) -> None:
    target = _parse_remind_target(command.args)
    if isinstance(target, ValueError):
        await message.answer(t(LANG, "remind_usage"))
        return
    if await sheets.is_closed():
        await message.answer(t(LANG, "remind_closed"))
        return
    students = await sheets.get_students_without_section(target)
    label = _target_label(target)
    if not students:
        await message.answer(t(LANG, "remind_nobody", target=label))
        return
    await message.answer(
        t(LANG, "remind_confirm", n=len(students), target=label),
        reply_markup=remind_confirm_kb("all" if target is None else str(target)),
    )


@router.callback_query(F.data == CB_REMIND_NO)
async def remind_no(callback: CallbackQuery) -> None:
    await callback.answer()
    if isinstance(callback.message, Message):
        await callback.message.edit_text(t(LANG, "remind_cancelled"))


@router.callback_query(F.data.startswith(f"{CB_REMIND_YES}:"))
async def remind_yes(callback: CallbackQuery) -> None:
    target = _parse_remind_target(callback.data.split(":", 2)[2])
    await callback.answer()
    message = callback.message if isinstance(callback.message, Message) else None
    if isinstance(target, ValueError) or message is None:
        return
    if await sheets.is_closed():
        await message.edit_text(t(LANG, "remind_closed"))
        return

    # Список берём заново: с момента подтверждения кто-то мог уже записаться.
    students = await sheets.get_students_without_section(target)
    await message.edit_text(t(LANG, "remind_sending", n=len(students)))

    sent = 0
    failed = 0
    for student in students:
        try:
            await callback.bot.send_message(student.tg_id, t(student.lang, "remind"))
            sent += 1
        except TelegramForbiddenError:
            # Ученик заблокировал бота.
            failed += 1
        except TelegramBadRequest as exc:
            # Например, «chat not found» — пользователь удалил аккаунт.
            failed += 1
            logger.warning("Не удалось отправить напоминание %s: %s", student.tg_id, exc)
        await asyncio.sleep(REMIND_DELAY_SECONDS)

    logger.info("Админ %s разослал напоминание (%s): отправлено %s, не доставлено %s",
                callback.from_user.id, _target_label(target), sent, failed)
    await message.edit_text(t(LANG, "remind_report", sent=sent, failed=failed))


# --------------------------------------------------------------------------- #
# /reload
# --------------------------------------------------------------------------- #


@router.message(Command("reload"))
async def cmd_reload(message: Message) -> None:
    await sheets.reload()
    logger.info(
        "Админ %s перечитал таблицу и обновил лист «%s»", message.from_user.id, sheets.ANALYTICS_SHEET
    )
    await message.answer(t(LANG, "reload_done"))
